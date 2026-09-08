"""输入解析：text / PDF / Excel → 统一文本。

PDF 用 pdfplumber（纯 Python，py3.14 友好）；Excel 用 openpyxl。
"""

from io import BytesIO


def parse_text(raw: str | bytes) -> str:
    if isinstance(raw, bytes):
        return raw.decode("utf-8", "ignore")
    return raw


def parse_pdf(file_bytes: bytes) -> str:
    import pdfplumber

    text_parts: list[str] = []
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t:
                text_parts.append(t)
    return "\n".join(text_parts).strip()


def parse_excel(file_bytes: bytes) -> list[str]:
    """每个数据行拼成一段文本（用首行做表头）。返回多条（一个客户一段）。"""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    texts: list[str] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            if not any(cells):
                continue
            parts = [f"{h}：{v}" for h, v in zip(header, cells) if v]
            if parts:
                texts.append("\n".join(parts))
    return texts
